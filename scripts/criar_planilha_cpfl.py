# Script para criar planilha de importação Ancoro a partir dos requisitos CPFL
# Projeto: Billing Risk - Melhorias em Faturamento

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
import os

def gerar_short_desc(frente, requisito):
    """Gera descrição curta (max 50 chars)"""
    palavras = requisito.split()[:6]
    desc = ' '.join(palavras)
    if len(desc) > 47:
        desc = desc[:47] + '...'
    return desc

def determinar_modulo(macro_requisito, frente):
    """Determina módulo SAP baseado no contexto"""
    frente_lower = frente.lower() if frente else ''
    macro_lower = macro_requisito.lower() if macro_requisito else ''

    if 'ccs' in macro_lower or 'billing' in frente_lower or 'faturamento' in frente_lower:
        return 'ISU-BILLING'
    elif 'crm' in macro_lower:
        return 'CRM'
    elif 'bpem' in macro_lower or 'medidor' in frente_lower:
        return 'ISU-BPEM'
    elif 'ide' in macro_lower or 'cadastro' in frente_lower:
        return 'ISU-IDE'
    else:
        return 'ISU-BILLING'

def extrair_who(requisito):
    """Extrai quem executa do texto"""
    requisito_lower = requisito.lower()
    if 'usuário' in requisito_lower or 'usuario' in requisito_lower:
        return 'Usuário do sistema'
    elif 'cliente' in requisito_lower:
        return 'Equipe de Faturamento'
    elif 'sistema' in requisito_lower:
        return 'Sistema SAP CCS'
    else:
        return 'Equipe de Faturamento'

def extrair_when(requisito, frente):
    """Extrai quando ocorre do texto"""
    requisito_lower = requisito.lower()
    if 'mensal' in requisito_lower:
        return 'Mensalmente'
    elif 'diário' in requisito_lower or 'diario' in requisito_lower:
        return 'Diariamente'
    elif 'automático' in requisito_lower or 'automatico' in requisito_lower:
        return 'Automaticamente durante faturamento'
    elif 'período de teste' in requisito_lower or 'periodo de teste' in requisito_lower:
        return 'Durante período de teste'
    else:
        return 'Durante ciclo de faturamento'

def extrair_where(macro_requisito):
    """Extrai onde acontece"""
    if 'CCS' in macro_requisito:
        return 'SAP CCS - Sistema de Faturamento'
    elif 'CRM' in macro_requisito:
        return 'SAP CRM - Gestão de Clientes'
    elif 'BW' in macro_requisito:
        return 'SAP BW - Business Warehouse'
    else:
        return 'SAP CCS - Sistema de Faturamento'

def gerar_how_today(requisito):
    """Gera descrição do processo atual (As-Is)"""
    req_lower = requisito.lower()
    if 'criar' in req_lower:
        return 'Funcionalidade não existe atualmente, requer implementação nova'
    elif 'ajustar' in req_lower or 'alterar' in req_lower:
        return 'Processo atual não atende todas as regras regulatórias, necessita ajuste'
    elif 'configurar' in req_lower:
        return 'Configuração atual incompleta ou inexistente no sistema'
    else:
        return 'Processo atual manual ou inexistente no sistema'

def gerar_why(frente, elegibilidade):
    """Gera justificativa do requisito"""
    frente_lower = frente.lower() if frente else ''

    if 'ren 901' in frente_lower:
        return 'Atendimento à Resolução Normativa 901 ANEEL para crédito de subvenção rural'
    elif 'liminar' in frente_lower:
        return 'Cumprimento de decisões judiciais sobre tributação ICMS'
    elif 'tusd' in frente_lower:
        return 'Aplicação correta dos descontos TUSD conforme regras CCEE'
    elif 'ape pie' in frente_lower:
        return 'Cálculo correto de autoprodução de energia para clientes livres'
    elif 'permission' in frente_lower:
        return 'Adequação ao novo modelo de suprimento parcial de permissionárias'
    elif 'demanda' in frente_lower:
        return 'Aplicação correta das regras de demanda para clientes rurais e sazonais'
    elif 'vencimento' in frente_lower:
        return 'Flexibilização de opções de pagamento para clientes'
    elif 'medidor' in frente_lower:
        return 'Garantir consistência no faturamento após movimentação de medidores'
    else:
        return 'Mitigação de riscos regulatórios e adequação às normas ANEEL'

def gerar_how_much(elegibilidade, frente):
    """Gera impacto do requisito"""
    if elegibilidade == 'Imprescindível':
        return 'Alto impacto - requisito regulatório obrigatório'
    elif elegibilidade == 'Importante':
        return 'Médio impacto - melhoria operacional'
    else:
        return 'Baixo impacto - otimização'


def main():
    # Caminhos
    input_path = 'C:/Users/rafae/OneDrive/Documentos/Profissional-Rafael/Documentos/CPFL/Anexo 1 - Matriz de Conformidade-v6.2.xlsx'
    output_path = 'c:/Users/rafae/OneDrive/Documentos/ancoro/requisitos_cpfl_billing_risk.xlsx'

    # Ler planilha CPFL
    print("Lendo planilha CPFL...")
    wb_cpfl = openpyxl.load_workbook(input_path)
    ws_cpfl = wb_cpfl['Requisitos']

    # Extrair requisitos funcionais
    requisitos = []
    for row in ws_cpfl.iter_rows(min_row=7, values_only=True):
        if row[0] and row[1] == 'Funcional':
            req = {
                'id_cpfl': str(row[0]).strip() if row[0] else '',
                'categoria': str(row[1]).strip() if row[1] else '',
                'secao': str(row[2]).strip() if row[2] else '',
                'frente': str(row[3]).strip() if row[3] else '',
                'agrupador': str(row[4]).strip() if row[4] else '',
                'macro_requisito': str(row[5]).strip() if row[5] else '',
                'requisito': str(row[6]).strip() if row[6] else '',
                'elegibilidade': str(row[7]).strip() if row[7] else '',
            }
            if req['requisito']:
                requisitos.append(req)

    print(f"Total de requisitos funcionais extraídos: {len(requisitos)}")

    # Criar planilha Ancoro
    print("Criando planilha para Ancoro...")
    wb_ancoro = Workbook()
    ws = wb_ancoro.active
    ws.title = "Requisitos"

    # Headers
    headers = ['reqId', 'shortDesc', 'module', 'what', 'why', 'who', 'when', 'where', 'howToday', 'howMuch', 'dependsOn', 'providesFor', 'status', 'observations']

    # Estilos
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # Escrever headers
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border

    # Popular dados
    for idx, req in enumerate(requisitos, 1):
        row_num = idx + 1
        req_id = f"REQ-{idx:03d}"

        requisito_texto = req['requisito']
        frente = req['frente']
        macro = req['macro_requisito']
        elegibilidade = req['elegibilidade']

        who = extrair_who(requisito_texto)
        when = extrair_when(requisito_texto, frente)
        where = extrair_where(macro)
        how_today = gerar_how_today(requisito_texto)
        why = gerar_why(frente, elegibilidade)
        how_much = gerar_how_much(elegibilidade, frente)

        dados = [
            req_id,
            gerar_short_desc(frente, requisito_texto),
            determinar_modulo(macro, frente),
            requisito_texto,
            why,
            who,
            when,
            where,
            how_today,
            how_much,
            '',
            '',
            'PENDING',
            f"Frente: {frente} | ID CPFL: {req['id_cpfl']} | {elegibilidade}"
        ]

        for col, valor in enumerate(dados, 1):
            cell = ws.cell(row=row_num, column=col, value=valor)
            cell.border = thin_border
            cell.alignment = Alignment(vertical='top', wrap_text=True)

    # Ajustar larguras
    column_widths = {
        'A': 10, 'B': 35, 'C': 15, 'D': 80, 'E': 60,
        'F': 25, 'G': 35, 'H': 35, 'I': 60, 'J': 40,
        'K': 15, 'L': 15, 'M': 12, 'N': 60
    }

    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width

    ws.freeze_panes = 'A2'

    # Salvar
    wb_ancoro.save(output_path)

    print(f"\nPlanilha criada com sucesso!")
    print(f"Arquivo: {output_path}")
    print(f"Total de requisitos: {len(requisitos)}")
    print("\nResumo por frente:")

    frentes = {}
    for req in requisitos:
        frente = req['frente']
        frentes[frente] = frentes.get(frente, 0) + 1

    for frente, count in sorted(frentes.items()):
        print(f"  - {frente}: {count}")


if __name__ == '__main__':
    main()
